"""
Standalone script to extract writing patterns from real כתב תביעה DOCX files.

Usage:
    python extract_firm_patterns.py /path/to/docx/folder

Reads 5-6 best כתב תביעה DOCX files, sends extracted text to Claude API,
and outputs firm_patterns.json with identified writing patterns.

This script is NOT imported by app.py — it's a one-time tool.
"""

import json
import os
import sys
import glob

from docx import Document
import anthropic


def extract_text_from_docx(filepath):
    """Extract all text from a DOCX file."""
    doc = Document(filepath)
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_firm_patterns.py /path/to/docx/folder")
        print("  Reads DOCX files from the folder and extracts firm writing patterns.")
        sys.exit(1)

    folder = sys.argv[1]
    if not os.path.isdir(folder):
        print(f"Error: {folder} is not a directory")
        sys.exit(1)

    # Find all DOCX files
    docx_files = glob.glob(os.path.join(folder, "*.docx"))
    if not docx_files:
        print(f"No DOCX files found in {folder}")
        sys.exit(1)

    # Sort by size (larger files tend to be more complete claims) and take top 6
    docx_files.sort(key=lambda f: os.path.getsize(f), reverse=True)
    selected = docx_files[:6]
    print(f"Found {len(docx_files)} DOCX files, using top {len(selected)}:")
    for f in selected:
        size_kb = os.path.getsize(f) / 1024
        print(f"  - {os.path.basename(f)} ({size_kb:.1f} KB)")

    # Extract text from each file
    all_texts = []
    for filepath in selected:
        try:
            text = extract_text_from_docx(filepath)
            if len(text) > 500:  # Only use substantial documents
                all_texts.append({
                    "filename": os.path.basename(filepath),
                    "text": text[:8000],  # Limit per doc to fit in context
                })
                print(f"  Extracted {len(text)} chars from {os.path.basename(filepath)}")
        except Exception as e:
            print(f"  Error reading {os.path.basename(filepath)}: {e}")

    if not all_texts:
        print("No usable text extracted from DOCX files")
        sys.exit(1)

    # Build prompt for Claude
    docs_text = ""
    for doc in all_texts:
        docs_text += f"\n\n--- Document: {doc['filename']} ---\n{doc['text']}"

    analysis_prompt = f"""Analyze the following Israeli labor law כתב תביעה documents from the Levin Telraz law firm.
Identify recurring writing patterns across the documents.

{docs_text}

Return a JSON object with this exact structure:
{{
  "patterns": {{
    "opening_phrases": ["list of common opening paragraph phrases used across documents"],
    "party_description_patterns": ["templates for describing plaintiff and defendant, using {{pronoun}}, {{name}}, {{id}} placeholders"],
    "claim_intro_phrases": {{
      "severance": "typical intro phrase for severance claims",
      "overtime": "typical intro phrase for overtime claims",
      "pension": "typical intro phrase for pension claims",
      "vacation": "typical intro phrase for vacation claims",
      "recuperation": "typical intro phrase for recuperation claims",
      "holidays": "typical intro phrase for holiday pay claims",
      "salary_delay": "typical intro phrase for salary delay claims",
      "deductions": "typical intro phrase for unlawful deductions claims",
      "emotional": "typical intro phrase for emotional distress claims",
      "prior_notice": "typical intro phrase for prior notice claims",
      "documents": "typical intro phrase for document delivery claims"
    }},
    "citation_patterns": ["common patterns for citing laws and regulations"],
    "narrative_connectors": ["transition words/phrases used between sections"],
    "closing_phrases": ["common closing/summary phrases"],
    "request_template": "the standard request paragraph template used when asking the court to award amounts"
  }}
}}

Important:
- Use {{placeholder}} syntax for variable parts (names, amounts, dates)
- Identify the FIRM'S specific style, not generic legal Hebrew
- Include gender-neutral placeholders where the firm uses gendered forms
- Return ONLY valid JSON, no markdown fences or explanation"""

    # Call Claude API
    api_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        print("Error: ANTHROPIC_API_KEY environment variable not set")
        sys.exit(1)

    client = anthropic.Anthropic(api_key=api_key)

    print("\nSending to Claude API for pattern analysis...")
    try:
        message = client.messages.create(
            model="claude-sonnet-4-5-20250929",
            max_tokens=4000,
            messages=[{"role": "user", "content": analysis_prompt}],
        )
        response_text = message.content[0].text.strip()

        # Parse and validate JSON
        if response_text.startswith("```"):
            lines = response_text.split("\n")
            json_lines = []
            in_fence = False
            for line in lines:
                if line.strip().startswith("```") and not in_fence:
                    in_fence = True
                    continue
                elif line.strip() == "```" and in_fence:
                    break
                elif in_fence:
                    json_lines.append(line)
            response_text = "\n".join(json_lines)

        patterns = json.loads(response_text)

        # Write output
        output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "firm_patterns.json")
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(patterns, f, ensure_ascii=False, indent=2)

        print(f"\nSuccess! Patterns written to {output_path}")
        print(f"Keys found: {list(patterns.get('patterns', {}).keys())}")

    except json.JSONDecodeError as e:
        print(f"Error: Claude returned invalid JSON: {e}")
        print(f"Raw response: {response_text[:500]}")
        sys.exit(1)
    except Exception as e:
        print(f"Error calling Claude API: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
